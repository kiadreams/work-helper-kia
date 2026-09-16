from pathlib import Path
import docx

from src.infrastructure.protocol_templates.template_render import (
    create_headline_template,
)


def test_create_headline_template(tmp_path: Path, mock_headline_data):
    project_root = Path(__file__).parents[4]
    template_file_path = (
            project_root
            / "src"
            / "infrastructure"
            / "protocol_templates"
            / "protocol_headline.docx"
    )
    target_file_path = tmp_path / "output_result.docx"

    assert template_file_path.exists(), f"No actual template found at the path: {template_file_path}"

    create_headline_template(template_file_path, target_file_path, mock_headline_data)

    assert target_file_path.exists(), "The final file was not created."

    result_doc = docx.Document(target_file_path.as_posix())
    text_pieces = [p.text for p in result_doc.paragraphs]
    # 2. Добавляем сбор текста из ВСЕХ таблиц документа
    for table in result_doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text_pieces.append(paragraph.text)

    # Объединяем всё в одну большую строку
    full_text = "".join(text_pieces)

    assert mock_headline_data["accession"] in full_text
    assert mock_headline_data["substation"] in full_text
    assert mock_headline_data["enterprise_name"] in full_text
    assert mock_headline_data["group_of_substation"] in full_text
    assert mock_headline_data["device_type_and_number"] in full_text
    assert mock_headline_data["creation_date"] in full_text
