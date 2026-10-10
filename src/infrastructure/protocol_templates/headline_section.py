from typing import TYPE_CHECKING

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from .styles import TextStyle, ParagraphStyle, CellStyle

if TYPE_CHECKING:
    from .protocols.headline_data_protocol import HeadlineDataProtocol
    from docx.table import _Cell


class HeadlineSection:
    text_style = TextStyle(text_align=WD_ALIGN_PARAGRAPH.CENTER, is_bold=True)
    paragraph_style = ParagraphStyle()
    cell_style = CellStyle()

    def __init__(self, headline_data: HeadlineDataProtocol) -> None:
        self._headline_data = headline_data
        self._sub_doc = Document()
        self._section = self._sub_doc.sections[0]
        self._section.top_margin = Cm(1.5)
        self._section.bottom_margin = Cm(1.5)
        self._section.left_margin = Cm(2)
        self._section.right_margin = Cm(1.5)

    def build(self):
        table = self._sub_doc.add_table(6, 2)
        for row in table.rows:
            for cell in row.cells:
                self._config_and_write_cell(
                    cell, "1111", self.text_style, self.paragraph_style, self.cell_style
                )

    # def inject_into(self, target_doc: Document):
    #     """Универсальный метод переноса содержимого в главный документ"""
    #     # Переносим параграфы
    #     for paragraph in self._sub_doc.paragraphs:
    #         new_p = target_doc.add_paragraph(paragraph.text, style=paragraph.style)
    #         new_p.alignment = paragraph.alignment
    #         new_p.paragraph_format.space_before = paragraph.paragraph_format.space_before
    #         new_p.paragraph_format.space_after = paragraph.paragraph_format.space_after
    #         new_p.paragraph_format.left_indent = paragraph.paragraph_format.left_indent
    #
    #         for run in paragraph.runs:
    #             new_run = new_p.add_run(run.text)
    #             new_run.bold = run.bold
    #             new_run.italic = run.italic
    #             if run.font.size:
    #                 new_run.font.size = run.font.size
    #
    #     # Переносим таблицы
    #     for table in self._sub_doc.tables:
    #         target_doc.element.body.append(table._element)

    def save_to(self, path: str):
        self._sub_doc.save(path)

    @staticmethod
    def _config_and_write_cell(
        cell: _Cell,
        text: str,
        text_style: TextStyle,
        paragraph_style: ParagraphStyle,
        cell_style: CellStyle,
    ) -> None:
        cell.vertical_alignment = cell_style.vertical_alignment
        p = cell.paragraphs[0]
        p.text = ""
        p.paragraph_format.space_after = paragraph_style.space_after
        p.paragraph_format.space_before = paragraph_style.space_before
        p.alignment = text_style.text_align

        run = p.add_run(text)
        run.bold = text_style.is_bold
        run.font.name = text_style.font_name
        run.font.size = text_style.font_size


if __name__ == "__main__":

    class ExampleData:
        def __init__(self):
            self._title = "Загловок"

        @property
        def title(self) -> str:
            return self._title

    headline = HeadlineSection(ExampleData())
    headline.build()
    headline.save_to("example.docx")
